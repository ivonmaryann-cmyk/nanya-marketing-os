from __future__ import annotations

import json
import math
import re
import uuid
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Mapping

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from werkzeug.utils import secure_filename

from .db import append_job_log, create_job, db_cursor, get_job, update_job_status
from .job_control import launch_job_process
from .paths import JOBS_DIR

FEATURE = "shennan_settlement"
RULE_VERSION = "深南结算规则 v1"
WAREHOUSES = ("0079", "0079B", "0082", "0082B", "0087", "0087B")
SEED_CONFIG = (("103673", "深南电路", "0079"), ("103673", "深南电路", "0079B"),
               ("104299", "无锡深南", "0082"), ("104299", "无锡深南", "0082B"),
               ("104370", "南通深南", "0087"), ("104370", "南通深南", "0087B"))
REASONS = ("未找到该客户物料代码对应的未结 341 订单", "对应订单未交数量已用完", "对应厂内料号无可用库存", "订单未交量与库存组合不足", "客户没有配置可用仓库")


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _id() -> str:
    return str(uuid.uuid4())


def init_schema() -> None:
    path = Path(__file__).parent.parent / "migrations/automation/postgresql/0015_shennan_settlement.sql"
    with db_cursor() as conn:
        for statement in path.read_text(encoding="utf-8").split(";"):
            if statement.strip():
                conn.execute(statement)
        for customer, name, warehouse in SEED_CONFIG:
            conn.execute("""INSERT INTO shennan_settlement_customer_warehouses
                (customer_code,customer_name,warehouse_code,updated_by,updated_at) VALUES (?,?,?,?,?)
                ON CONFLICT(customer_code,warehouse_code) DO NOTHING""", (customer, name, warehouse, "标准初始化", _now()))


def list_config() -> list[dict[str, Any]]:
    with db_cursor() as conn:
        return [dict(row) for row in conn.execute("SELECT * FROM shennan_settlement_customer_warehouses ORDER BY customer_code,warehouse_code").fetchall()]


def configured_warehouse_codes() -> tuple[str, ...]:
    with db_cursor() as conn:
        rows = conn.execute("SELECT DISTINCT warehouse_code FROM shennan_settlement_customer_warehouses WHERE enabled=1 ORDER BY warehouse_code").fetchall()
    return tuple(row["warehouse_code"] for row in rows)


def save_config(customer_code: str, customer_name: str, warehouse_code: str, enabled: bool, actor: str) -> None:
    customer_code, customer_name, warehouse_code = customer_code.strip(), customer_name.strip(), warehouse_code.strip().upper()
    if not customer_code.isdigit() or not customer_name or not warehouse_code:
        raise ValueError("客户编号、客户名称和仓库编码均为必填项")
    with db_cursor() as conn:
        conn.execute("""INSERT INTO shennan_settlement_customer_warehouses
          (customer_code,customer_name,warehouse_code,enabled,updated_by,updated_at) VALUES (?,?,?,?,?,?)
          ON CONFLICT(customer_code,warehouse_code) DO UPDATE SET customer_name=excluded.customer_name,
          enabled=excluded.enabled,revision=shennan_settlement_customer_warehouses.revision+1,
          updated_by=excluded.updated_by,updated_at=excluded.updated_at""",
          (customer_code, customer_name, warehouse_code, int(enabled), actor, _now()))


def list_batches(employee_id: str) -> list[dict[str, Any]]:
    with db_cursor() as conn:
        return [dict(row) for row in conn.execute("SELECT * FROM shennan_settlement_batches WHERE employee_id=? ORDER BY created_at DESC", (employee_id,)).fetchall()]


def get_batch(batch_id: str, employee_id: str | None = None) -> dict[str, Any] | None:
    query, params = "SELECT * FROM shennan_settlement_batches WHERE id=?", [batch_id]
    if employee_id:
        query += " AND employee_id=?"; params.append(employee_id)
    with db_cursor() as conn:
        row = conn.execute(query, params).fetchone()
        return dict(row) if row else None


def batch_view(batch_id: str, employee_id: str) -> dict[str, Any] | None:
    batch = get_batch(batch_id, employee_id)
    if not batch:
        return None
    with db_cursor() as conn:
        for key, table in (("details", "shennan_settlement_details"), ("unmatched", "shennan_settlement_unmatched")):
            batch[key] = [dict(row) for row in conn.execute(f"SELECT * FROM {table} WHERE batch_id=? ORDER BY block_no,source_row", (batch_id,)).fetchall()]
    return batch


def normalize_inventory_uploads(inventory_files: Mapping[str, Any] | Iterable[Any], required_warehouses: Iterable[str]) -> dict[str, Any]:
    """Map a multi-file browser upload to configured warehouse codes by filename."""
    required = tuple(required_warehouses)
    if isinstance(inventory_files, Mapping):
        files = list(inventory_files.items())
    else:
        files = []
        for file_obj in inventory_files:
            filename = str(getattr(file_obj, "filename", "") or "")
            files.append((Path(filename).stem.upper(), file_obj))
    resolved: dict[str, Any] = {}
    invalid: list[str] = []
    unknown: list[str] = []
    duplicate: list[str] = []
    for warehouse, file_obj in files:
        filename = str(getattr(file_obj, "filename", "") or "")
        if not filename:
            continue
        if Path(filename).suffix.lower() not in {".xls", ".xlsx"}:
            invalid.append(filename)
            continue
        warehouse = str(warehouse).strip().upper()
        if warehouse not in required:
            unknown.append(filename)
        elif warehouse in resolved:
            duplicate.append(warehouse)
        else:
            resolved[warehouse] = file_obj
    if invalid:
        raise ValueError("库存仅支持 .xls / .xlsx：" + "、".join(invalid))
    if unknown:
        raise ValueError("库存文件名未对应启用仓库：" + "、".join(unknown))
    if duplicate:
        raise ValueError("重复上传仓库库存表：" + "、".join(duplicate))
    missing = [code for code in required if code not in resolved]
    if missing:
        raise ValueError("请上传库存表：" + "、".join(missing))
    return resolved


def queue_batch(employee_id: str, consumption_file, order_file, inventory_files: Mapping[str, Any] | Iterable[Any]) -> tuple[str, int]:
    if not consumption_file or not consumption_file.filename or not order_file or not order_file.filename:
        raise ValueError("请上传客户消耗表和 341 未结订单表")
    required_warehouses = configured_warehouse_codes()
    inventory_files = normalize_inventory_uploads(inventory_files, required_warehouses)
    invalid = [f.filename for f in [consumption_file, order_file] if Path(f.filename).suffix.lower() not in {".xls", ".xlsx"}]
    if invalid:
        raise ValueError("仅支持 .xls / .xlsx：" + "、".join(invalid))
    batch_id, stamp = _id(), datetime.now().strftime("%Y%m%d_%H%M%S")
    folder = JOBS_DIR / employee_id / f"{stamp}_shennan_settlement"
    folder.mkdir(parents=True, exist_ok=True)
    def save(file_obj, label):
        suffix = Path(file_obj.filename).suffix.lower(); target = folder / f"{label}_{secure_filename(Path(file_obj.filename).stem) or label}{suffix}"
        file_obj.save(target); return target
    manifest = {"batch_id": batch_id, "consumption": str(save(consumption_file, "consumption")), "orders": str(save(order_file, "orders")), "inventory": {code: str(save(inventory_files[code], f"inventory_{code}")) for code in required_warehouses}}
    manifest_path = folder / "shennan_settlement_input.json"; manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    now = _now()
    with db_cursor() as conn:
        conn.execute("INSERT INTO shennan_settlement_batches(id,employee_id,status,manifest_path,created_at,updated_at) VALUES (?,?, 'queued',?,?,?)", (batch_id, employee_id, str(manifest_path), now, now))
        conn.execute("INSERT INTO shennan_settlement_events VALUES (?,?,?,?,?,?)", (_id(), batch_id, "created", employee_id, "已创建草稿结算批次", now))
    job_id = create_job(employee_id, f"深南结算：{consumption_file.filename}", str(manifest_path), RULE_VERSION, feature=FEATURE)
    with db_cursor() as conn: conn.execute("UPDATE shennan_settlement_batches SET job_id=?,updated_at=? WHERE id=?", (job_id, _now(), batch_id))
    launch_job_process(job_id, FEATURE, employee_id)
    return batch_id, job_id


def queue_recalculate(batch_id: str, employee_id: str) -> int:
    batch = get_batch(batch_id, employee_id)
    if not batch or batch["status"] not in {"draft", "failed"}: raise ValueError("只有草稿或失败批次可以重新计算")
    job_id = create_job(employee_id, f"深南结算重新计算：{batch_id[:8]}", batch["manifest_path"], RULE_VERSION, feature=FEATURE)
    with db_cursor() as conn: conn.execute("UPDATE shennan_settlement_batches SET job_id=?,status='queued',updated_at=? WHERE id=?", (job_id, _now(), batch_id))
    launch_job_process(job_id, FEATURE, employee_id); return job_id


def run_shennan_settlement_job(job_id: int, employee_id: str) -> None:
    job = get_job(job_id)
    if not job or job["employee_id"] != employee_id: return
    update_job_status(job_id, status="running", log_text="")
    try:
        manifest = json.loads(Path(job["stored_input_path"]).read_text(encoding="utf-8")); batch_id = manifest["batch_id"]
        append_job_log(job_id, "开始读取深南结算源文件。")
        result = calculate_batch(batch_id, manifest, employee_id)
        update_job_status(job_id, status="completed", stored_result_path=result["output_path"], success_count=result["detail_count"], fail_count=result["unmatched_count"], current_row=result["consumption_count"], total_rows=result["consumption_count"], completed=True)
        append_job_log(job_id, f"结算草稿完成：分摊 {result['detail_count']} 条，未匹配 {result['unmatched_count']} 条。")
    except Exception as exc:
        append_job_log(job_id, f"深南结算处理失败：{exc}")
        update_job_status(job_id, status="failed", error_message=str(exc), completed=True)
        try:
            with db_cursor() as conn: conn.execute("UPDATE shennan_settlement_batches SET status='failed',updated_at=? WHERE manifest_path=?", (_now(), job["stored_input_path"]))
        except Exception: pass
        raise


def calculate_batch(batch_id: str, manifest: dict[str, Any], actor: str) -> dict[str, Any]:
    consumptions = _read_consumptions(Path(manifest["consumption"])); orders = _read_orders(Path(manifest["orders"])); inventory = _read_inventory(manifest["inventory"])
    config = _enabled_config()
    with db_cursor() as conn:
        _clear_draft(conn, batch_id)
        _snapshot(conn, batch_id, manifest, consumptions, orders, inventory)
        confirmed_orders, confirmed_inventory = _confirmed_usage(conn)
        details, unmatched = _match(consumptions, orders, inventory, config, confirmed_orders, confirmed_inventory)
        _validate(consumptions, details, unmatched, config)
        _persist_result(conn, batch_id, details, unmatched)
    output_path = Path(manifest["consumption"]).parent / f"深南结算_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    write_workbook(details, unmatched, output_path)
    with db_cursor() as conn:
        conn.execute("UPDATE shennan_settlement_batches SET status='draft',output_path=?,updated_at=? WHERE id=?", (str(output_path), _now(), batch_id))
        conn.execute("INSERT INTO shennan_settlement_events VALUES (?,?,?,?,?,?)", (_id(), batch_id, "calculated", actor, "已生成草稿结算结果", _now()))
    return {"output_path": str(output_path), "detail_count": len(details), "unmatched_count": len(unmatched), "consumption_count": len(consumptions)}


def _text(value: Any) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)): return ""
    if isinstance(value, float) and value.is_integer(): return str(int(value))
    return str(value).strip()


def _number(value: Any, field: str) -> float:
    try:
        n = float(str(value).replace(",", "").strip())
    except (TypeError, ValueError): raise ValueError(f"{field} 不是有效数字")
    if not math.isfinite(n): raise ValueError(f"{field} 不是有效数字")
    return n


def _date(value: Any) -> str:
    text = _text(value)
    if not text: return ""
    # Inventory HTML exports use YY/MM/DD (for example 26/06/27).  Letting
    # pandas infer this format treats it as DD/MM/YY and reverses FIFO order.
    if re.fullmatch(r"\d{2}/\d{1,2}/\d{1,2}", text):
        return datetime.strptime(text, "%y/%m/%d").strftime("%Y-%m-%d")
    parsed = pd.to_datetime(value, errors="coerce")
    return parsed.strftime("%Y-%m-%d") if not pd.isna(parsed) else text


def _read_consumptions(path: Path) -> list[dict[str, Any]]:
    frame = pd.read_excel(path, sheet_name=0, header=None, dtype=object)
    rows, block = [], 0
    for idx, row in frame.iterrows():
        values = row.tolist()
        if not any(_text(value) for value in values): block += 1; continue
        if not block: block = 1
        if len(values) < 4: raise ValueError(f"客户消耗表第 {idx + 1} 行字段不足")
        customer, material = _text(values[0]), _text(values[1])
        quantity, price = _number(values[2], f"客户消耗表第 {idx + 1} 行消耗数量"), _number(values[3], f"客户消耗表第 {idx + 1} 行价格")
        if not customer or not material or quantity <= 0: raise ValueError(f"客户消耗表第 {idx + 1} 行客户、物料和正数数量为必填")
        rows.append({"id": _id(), "block_no": block, "source_row": idx + 1, "customer_code": customer, "customer_material_code": material, "quantity": quantity, "price": price})
    if not rows: raise ValueError("客户消耗表没有可处理数据")
    return rows


def _read_orders(path: Path) -> list[dict[str, Any]]:
    frame = pd.read_excel(path, sheet_name=0, header=2, dtype=object)
    required = {"单别单号", "项次", "订单日期", "客户编号", "客户产品编号", "品号", "未交数量"}
    if missing := required - set(frame.columns): raise ValueError("341订单表缺少字段：" + "、".join(sorted(missing)))
    rows=[]
    for _, r in frame.iterrows():
        no, customer, material, part = (_text(r["单别单号"]), _text(r["客户编号"]), _text(r["客户产品编号"]), _text(r["品号"]))
        if not no.startswith("341-") or not customer or not material or not part: continue
        try: quantity = _number(r["未交数量"], "订单未交数量")
        except ValueError: continue
        if quantity <= 0: continue
        rows.append({"id":_id(),"order_no":no,"line_no":_text(r["项次"]),"order_date":_date(r["订单日期"]),"customer_code":customer,"customer_material_code":material,"factory_part_no":part,"quantity":quantity})
    return rows


def _read_inventory(paths: dict[str, str]) -> list[dict[str, Any]]:
    rows=[]
    for warehouse, raw_path in paths.items():
        path=Path(raw_path)
        try:
            if path.read_bytes()[:20].lstrip().lower().startswith(b"<html"):
                frame=pd.read_html(path, encoding="utf-8")[0]; frame.columns=[_text(v) for v in frame.iloc[0].tolist()]; frame=frame.iloc[1:]
            else: frame=pd.read_excel(path, sheet_name=0, dtype=object)
        except Exception as exc: raise ValueError(f"库存文件 {warehouse} 无法读取：{exc}") from exc
        required={"料件编号","批号","库存数量","呆滞日期"}
        if missing := required-set(frame.columns): raise ValueError(f"库存文件 {warehouse} 缺少字段："+"、".join(sorted(missing)))
        for _, r in frame.iterrows():
            part, lot=_text(r["料件编号"]),_text(r["批号"])
            if not part or not lot: continue
            try: qty=_number(r["库存数量"], "库存数量")
            except ValueError: continue
            if qty<=0: continue
            rows.append({"id":_id(),"warehouse_code":warehouse,"factory_part_no":part,"lot_no":lot,"quantity":qty,"stagnation_date":_date(r["呆滞日期"])})
    return rows


def _enabled_config() -> dict[str, dict[str, Any]]:
    with db_cursor() as conn:
        rows = conn.execute("SELECT * FROM shennan_settlement_customer_warehouses WHERE enabled=1").fetchall()
    result: dict[str, dict[str, Any]] = {}
    for row in rows:
        item = result.setdefault(row["customer_code"], {"name": row["customer_name"], "warehouses": set()})
        item["warehouses"].add(row["warehouse_code"])
    return result


def _confirmed_usage(conn) -> tuple[dict[tuple[str, str], float], dict[tuple[str, str, str], float]]:
    orders, inventory = defaultdict(float), defaultdict(float)
    rows = conn.execute("""SELECT d.* FROM shennan_settlement_details d
        JOIN shennan_settlement_batches b ON b.id=d.batch_id WHERE b.status='confirmed'""").fetchall()
    for row in rows:
        orders[(row["order_no"], row["line_no"])] += float(row["quantity"])
        inventory[(row["warehouse_code"], row["factory_part_no"], row["lot_no"])] += float(row["quantity"])
    return orders, inventory


def _match(consumptions, orders, inventory, config, confirmed_orders, confirmed_inventory):
    orders_by_key=defaultdict(list); stock_by_part=defaultdict(list)
    for order in orders:
        order["remaining"] = max(0.0, order["quantity"] - confirmed_orders.get((order["order_no"], order["line_no"]), 0.0))
        orders_by_key[(order["customer_code"], order["customer_material_code"])].append(order)
    for value in orders_by_key.values(): value.sort(key=lambda r:(r["order_date"],r["order_no"],_sort_line(r["line_no"])))
    for stock in inventory:
        stock["remaining"] = max(0.0, stock["quantity"] - confirmed_inventory.get((stock["warehouse_code"], stock["factory_part_no"], stock["lot_no"]), 0.0))
        stock_by_part[stock["factory_part_no"]].append(stock)
    for value in stock_by_part.values(): value.sort(key=lambda r:(r["stagnation_date"] or "9999-12-31",r["warehouse_code"],r["lot_no"]))
    details=[]; unmatched=[]
    for consume in consumptions:
        remaining=consume["quantity"]; candidates=orders_by_key[(consume["customer_code"],consume["customer_material_code"])]
        info=config.get(consume["customer_code"])
        if not info:
            unmatched.append(_unmatched(consume,0,REASONS[4],"")); continue
        if not candidates:
            unmatched.append(_unmatched(consume,0,REASONS[0],info["name"])); continue
        had_open_order=False; had_stock=False
        for order in candidates:
            if remaining <= 1e-9: break
            if order["remaining"] <= 1e-9: continue
            had_open_order=True
            stocks=[x for x in stock_by_part[order["factory_part_no"]] if x["warehouse_code"] in info["warehouses"]]
            if any(x["remaining"] > 1e-9 for x in stocks): had_stock=True
            for stock in stocks:
                if remaining <= 1e-9 or order["remaining"] <= 1e-9: break
                if stock["remaining"] <= 1e-9: continue
                qty=min(remaining,order["remaining"],stock["remaining"])
                details.append({"id":_id(),"consumption_id":consume["id"],"block_no":consume["block_no"],"source_row":consume["source_row"],"customer_code":consume["customer_code"],"customer_name":info["name"],"category":"PP" if order["factory_part_no"].startswith("69") else "基板","order_no":order["order_no"],"line_no":order["line_no"],"order_date":order["order_date"],"customer_material_code":consume["customer_material_code"],"factory_part_no":order["factory_part_no"],"lot_no":stock["lot_no"],"warehouse_code":stock["warehouse_code"],"price":consume["price"],"quantity":qty})
                remaining-=qty; order["remaining"]-=qty; stock["remaining"]-=qty
        matched=consume["quantity"]-remaining
        if remaining > 1e-9:
            reason=REASONS[1] if not had_open_order else REASONS[2] if not had_stock else REASONS[3]
            unmatched.append(_unmatched(consume,matched,reason,info["name"]))
    return details, unmatched


def _sort_line(value: str):
    try: return (0,int(value))
    except ValueError: return (1,value)


def _unmatched(consumption, matched, reason, name):
    return {"id":_id(),"block_no":consumption["block_no"],"source_row":consumption["source_row"],"customer_name":name,"customer_code":consumption["customer_code"],"customer_material_code":consumption["customer_material_code"],"original_quantity":consumption["quantity"],"matched_quantity":matched,"unmatched_quantity":consumption["quantity"]-matched,"price":consumption["price"],"reason":reason}


def _validate(consumptions, details, unmatched, config):
    by_consume=defaultdict(float)
    for detail in details:
        if detail["quantity"] <= 0: raise ValueError("结算数量必须大于 0")
        if detail["warehouse_code"] not in config[detail["customer_code"]]["warehouses"]: raise ValueError("结算明细使用了未配置仓库")
        by_consume[detail["consumption_id"]] += detail["quantity"]
    unmatched_by_source={row["source_row"]:row for row in unmatched}
    for row in consumptions:
        missing=unmatched_by_source.get(row["source_row"],{}).get("unmatched_quantity",0)
        if abs(row["quantity"]-by_consume[row["id"]]-missing)>1e-6: raise ValueError("客户消耗总量校验失败")


def _clear_draft(conn, batch_id):
    for table in ("shennan_settlement_inputs","shennan_settlement_consumptions","shennan_settlement_orders","shennan_settlement_inventory","shennan_settlement_details","shennan_settlement_unmatched"):
        conn.execute(f"DELETE FROM {table} WHERE batch_id=?", (batch_id,))


def _snapshot(conn, batch_id, manifest, consumptions, orders, inventory):
    now=_now()
    inputs=[("consumption","",manifest["consumption"]),("orders","",manifest["orders"]),*[("inventory",code,path) for code,path in manifest["inventory"].items()]]
    for kind,warehouse,path in inputs: conn.execute("INSERT INTO shennan_settlement_inputs VALUES (?,?,?,?,?,?,?,?)", (_id(),batch_id,kind,warehouse,Path(path).name,path,0,now))
    conn.executemany("INSERT INTO shennan_settlement_consumptions VALUES (?,?,?,?,?,?,?,?,?,?)", [(r["id"],batch_id,r["block_no"],r["source_row"],r["customer_code"],r["customer_material_code"],r["quantity"],r["price"],0,r["quantity"]) for r in consumptions])
    conn.executemany("INSERT INTO shennan_settlement_orders VALUES (?,?,?,?,?,?,?,?,?)", [(r["id"],batch_id,r["order_no"],r["line_no"],r["order_date"],r["customer_code"],r["customer_material_code"],r["factory_part_no"],r["quantity"]) for r in orders])
    conn.executemany("INSERT INTO shennan_settlement_inventory VALUES (?,?,?,?,?,?,?)", [(r["id"],batch_id,r["warehouse_code"],r["factory_part_no"],r["lot_no"],r["quantity"],r["stagnation_date"]) for r in inventory])


def _persist_result(conn, batch_id, details, unmatched):
    conn.executemany("INSERT INTO shennan_settlement_details VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", [(r["id"],batch_id,r["consumption_id"],r["block_no"],r["source_row"],r["customer_code"],r["customer_name"],r["category"],r["order_no"],r["line_no"],r["order_date"],r["customer_material_code"],r["factory_part_no"],r["lot_no"],r["warehouse_code"],r["price"],r["quantity"]) for r in details])
    conn.executemany("INSERT INTO shennan_settlement_unmatched VALUES (?,?,?,?,?,?,?,?,?,?,?,?)", [(r["id"],batch_id,r["block_no"],r["source_row"],r["customer_name"],r["customer_code"],r["customer_material_code"],r["original_quantity"],r["matched_quantity"],r["unmatched_quantity"],r["price"],r["reason"]) for r in unmatched])
    for row in details: conn.execute("UPDATE shennan_settlement_consumptions SET matched_quantity=matched_quantity+?,remaining_quantity=remaining_quantity-? WHERE id=?", (row["quantity"],row["quantity"],row["consumption_id"]))


def confirm_batch(batch_id: str, employee_id: str) -> None:
    with db_cursor() as conn:
        if getattr(conn,"dialect","")=="postgresql": conn.execute("LOCK TABLE shennan_settlement_batches, shennan_settlement_details IN SHARE ROW EXCLUSIVE MODE")
        batch=conn.execute("SELECT * FROM shennan_settlement_batches WHERE id=? AND employee_id=?",(batch_id,employee_id)).fetchone()
        if not batch or batch["status"]!="draft": raise ValueError("仅可确认已计算完成的草稿批次")
        confirmed_orders,confirmed_inventory=_confirmed_usage(conn)
        orders={ (r["order_no"],r["line_no"]):float(r["unfulfilled_quantity"]) for r in conn.execute("SELECT * FROM shennan_settlement_orders WHERE batch_id=?",(batch_id,)).fetchall() }
        stock={ (r["warehouse_code"],r["factory_part_no"],r["lot_no"]):float(r["available_quantity"]) for r in conn.execute("SELECT * FROM shennan_settlement_inventory WHERE batch_id=?",(batch_id,)).fetchall() }
        used_orders,used_stock=defaultdict(float),defaultdict(float)
        for r in conn.execute("SELECT * FROM shennan_settlement_details WHERE batch_id=?",(batch_id,)).fetchall():
            qty=float(r["quantity"]); ok=(r["order_no"],r["line_no"]); sk=(r["warehouse_code"],r["factory_part_no"],r["lot_no"]); used_orders[ok]+=qty; used_stock[sk]+=qty
        if any(used_orders[key] > orders.get(key,0)-confirmed_orders[key]+1e-6 for key in used_orders) or any(used_stock[key] > stock.get(key,0)-confirmed_inventory[key]+1e-6 for key in used_stock): raise ValueError("订单或库存已被其他正式批次占用，请重新计算后再确认")
        now=_now(); conn.execute("UPDATE shennan_settlement_batches SET status='confirmed',confirmed_at=?,updated_at=? WHERE id=?",(now,now,batch_id)); conn.execute("INSERT INTO shennan_settlement_events VALUES (?,?,?,?,?,?)",(_id(),batch_id,"confirmed",employee_id,"已正式确认并占用订单与库存",now))


def cancel_confirmation(batch_id: str, employee_id: str) -> None:
    with db_cursor() as conn:
        row=conn.execute("SELECT * FROM shennan_settlement_batches WHERE id=? AND employee_id=?",(batch_id,employee_id)).fetchone()
        if not row or row["status"]!="confirmed": raise ValueError("仅可撤销已确认批次")
        now=_now(); conn.execute("UPDATE shennan_settlement_batches SET status='cancelled',cancelled_at=?,updated_at=? WHERE id=?",(now,now,batch_id)); conn.execute("INSERT INTO shennan_settlement_events VALUES (?,?,?,?,?,?)",(_id(),batch_id,"cancelled",employee_id,"已撤销确认并释放占用",now))


def write_workbook(details, unmatched, output_path: Path):
    workbook=Workbook(); workbook.remove(workbook.active); counters=defaultdict(int); groups=defaultdict(list)
    for row in details: groups[(row["block_no"],row["customer_code"],row["customer_name"],row["category"])].append(row)
    for _, rows in sorted(groups.items(),key=lambda item:(item[0][0],item[0][1],item[0][3])):
        sample=rows[0]; counter=(sample["customer_code"],sample["category"]); counters[counter]+=1; sheet=workbook.create_sheet(f"{sample['customer_name']}{sample['category']}{counters[counter]}"[:31]); sheet["A1"]="出货单号"; sheet["B1"]="客户编号"; sheet["A2"]=""; sheet["B2"]=sample["customer_code"]
        headers=["341订单号","订单项次","结算数量","批号","价格","仓库"]
        for index,value in enumerate(headers,1): sheet.cell(3,index,value)
        for row in rows: sheet.append([row["order_no"],row["line_no"],row["quantity"],row["lot_no"],row["price"],row["warehouse_code"]])
        _style(sheet,6); sheet.column_dimensions["E"].width=14
        for cell in sheet["E"][3:]: cell.number_format="0.0000"
    sheet=workbook.create_sheet("未匹配明细"); headers=["数据块","客户","客户编号","客户物料代码","源数据行","原消耗数量","已匹配数量","未匹配数量","价格","原因"]; sheet.append(headers)
    for row in unmatched: sheet.append([row["block_no"],row["customer_name"],row["customer_code"],row["customer_material_code"],row["source_row"],row["original_quantity"],row["matched_quantity"],row["unmatched_quantity"],row["price"],row["reason"]])
    _style(sheet,len(headers)); sheet.column_dimensions["J"].width=40
    for cell in sheet["I"][1:]: cell.number_format="0.0000"
    output_path.parent.mkdir(parents=True,exist_ok=True); workbook.save(output_path)


def _style(sheet, columns):
    border=Border(left=Side(style="thin"),right=Side(style="thin"),top=Side(style="thin"),bottom=Side(style="thin")); fill=PatternFill("solid",fgColor="D9EAF7")
    for row in sheet.iter_rows():
        for cell in row: cell.border=border; cell.alignment=Alignment(horizontal="center",vertical="center"); cell.font=Font(name="宋体",size=10)
    for cell in sheet[3] if sheet.title!="未匹配明细" else sheet[1]: cell.fill=fill; cell.font=Font(name="宋体",size=10,bold=True)
    for index in range(1,columns+1): sheet.column_dimensions[sheet.cell(1,index).column_letter].width=16
    sheet.freeze_panes="A4" if sheet.title!="未匹配明细" else "A2"
