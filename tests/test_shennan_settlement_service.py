from __future__ import annotations

import sqlite3
import tempfile
import unittest
from contextlib import contextmanager
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from fangzheng_web_app import shennan_settlement_service as service


@contextmanager
def _cursor(connection):
    yield connection


class ShennanSettlementServiceTests(unittest.TestCase):
    def setUp(self):
        self.connection = sqlite3.connect(":memory:")
        self.connection.row_factory = sqlite3.Row
        self.db_patch = patch.object(service, "db_cursor", lambda: _cursor(self.connection))
        self.db_patch.start()
        service.init_schema()
        self.batch_id = "batch-1"
        self.connection.execute("""INSERT INTO shennan_settlement_batches
            (id,employee_id,status,manifest_path,created_at,updated_at) VALUES (?,?, 'queued','',?,?)""",
            (self.batch_id, "u1", service._now(), service._now()))

    def tearDown(self):
        self.db_patch.stop()
        self.connection.close()

    def _files(self, directory: Path):
        consumption = directory / "consume.xlsx"
        wb = Workbook(); sheet = wb.active
        sheet.append(["103673", "C1", 12, 10])
        sheet.append(["103673", "C1", 3, 11])
        sheet.append([None, None, None, None])
        sheet.append(["103673", "C2", 5, 12])
        wb.save(consumption)

        orders = directory / "orders.xlsx"
        wb = Workbook(); sheet = wb.active
        sheet.append(["标题"]); sheet.append(["标题"])
        headers = ["单别单号", "项次", "订单日期", "客户编号", "客户产品编号", "品号", "未交数量"]
        sheet.append(headers)
        sheet.append(["341-2", 2, "2026-01-02", "103673", "C1", "6901", 10])
        sheet.append(["341-1", 1, "2026-01-01", "103673", "C1", "6901", 8])
        sheet.append(["341-3", 1, "2026-01-03", "103673", "C2", "7101", 5])
        wb.save(orders)

        inventory = {}
        for warehouse in service.WAREHOUSES:
            path = directory / f"{warehouse}.xls"
            rows = ""
            if warehouse == "0079":
                rows = "<tr><td>6901</td><td>x</td><td>x</td><td></td><td>L2</td><td>x</td><td>5</td><td></td><td></td><td>26/02/01</td></tr><tr><td>6901</td><td>x</td><td>x</td><td></td><td>L1</td><td>x</td><td>10</td><td></td><td></td><td>26/01/01</td></tr><tr><td>7101</td><td>x</td><td>x</td><td></td><td>B1</td><td>x</td><td>5</td><td></td><td></td><td>26/01/01</td></tr><tr><td>6901</td><td>x</td><td>x</td><td></td><td>ZERO</td><td>x</td><td>0</td><td></td><td></td><td>26/01/01</td></tr>"
            header = "<tr><td>料件编号</td><td>品名</td><td>规格</td><td>库位</td><td>批号</td><td>库存单位</td><td>库存数量</td><td>库存等级</td><td>外观编号</td><td>呆滞日期</td></tr>"
            path.write_text(f"<html><table>{header}{rows}</table></html>", encoding="utf-8")
            inventory[warehouse] = str(path)
        return {"batch_id": self.batch_id, "consumption": str(consumption), "orders": str(orders), "inventory": inventory}

    def test_calculates_orders_and_lots_in_required_order_and_exports(self):
        with tempfile.TemporaryDirectory() as directory:
            result = service.calculate_batch(self.batch_id, self._files(Path(directory)), "u1")
            details = self.connection.execute("SELECT * FROM shennan_settlement_details ORDER BY block_no,source_row,order_no,lot_no").fetchall()
            self.assertEqual(result["detail_count"], 5)
            self.assertEqual([(r["order_no"], r["lot_no"], r["quantity"]) for r in details[:3]], [("341-1", "L1", 8.0), ("341-2", "L1", 2.0), ("341-2", "L2", 2.0)])
            self.assertTrue(all(r["warehouse_code"] == "0079" for r in details))
            self.assertNotIn("ZERO", [r["lot_no"] for r in details])
            workbook = load_workbook(result["output_path"], data_only=True)
            self.assertIn("深南电路PP1", workbook.sheetnames)
            self.assertIn("深南电路基板1", workbook.sheetnames)
            self.assertIn("未匹配明细", workbook.sheetnames)
            self.assertEqual(workbook["深南电路PP1"]["A3"].value, "341订单号")
            self.assertEqual(workbook["深南电路PP1"]["F4"].value, "0079")

    def test_confirmation_locks_and_cancellation_releases_usage(self):
        with tempfile.TemporaryDirectory() as directory:
            service.calculate_batch(self.batch_id, self._files(Path(directory)), "u1")
            service.confirm_batch(self.batch_id, "u1")
            self.assertEqual(service.get_batch(self.batch_id, "u1")["status"], "confirmed")
            service.cancel_confirmation(self.batch_id, "u1")
            self.assertEqual(service.get_batch(self.batch_id, "u1")["status"], "cancelled")

    def test_reports_missing_customer_configuration(self):
        consumption = [{"id":"c1", "block_no":1, "source_row":2, "customer_code":"999", "customer_material_code":"x", "quantity":1, "price":1}]
        details, unmatched = service._match(consumption, [], [], {}, {}, {})
        self.assertEqual(details, [])
        self.assertEqual(unmatched[0]["reason"], "客户没有配置可用仓库")

    def test_classifies_all_order_and_stock_shortages(self):
        consumption = [{"id":"c1", "block_no":1, "source_row":2, "customer_code":"1", "customer_material_code":"x", "quantity":2, "price":1}]
        config = {"1": {"name":"客户", "warehouses":{"W"}}}
        self.assertEqual(service._match(consumption, [], [], config, {}, {})[1][0]["reason"], "未找到该客户物料代码对应的未结 341 订单")
        exhausted = [{"order_no":"341-1", "line_no":"1", "order_date":"2026-01-01", "customer_code":"1", "customer_material_code":"x", "factory_part_no":"69", "quantity":0}]
        self.assertEqual(service._match(consumption, exhausted, [], config, {}, {})[1][0]["reason"], "对应订单未交数量已用完")
        order = [{**exhausted[0], "quantity":2}]
        self.assertEqual(service._match(consumption, order, [], config, {}, {})[1][0]["reason"], "对应厂内料号无可用库存")
        stock = [{"warehouse_code":"W", "factory_part_no":"69", "lot_no":"L", "quantity":1, "stagnation_date":"2026-01-01"}]
        self.assertEqual(service._match(consumption, order, stock, config, {}, {})[1][0]["reason"], "订单未交量与库存组合不足")

    def test_parses_inventory_yy_mm_dd_and_consumes_the_oldest_lot_first(self):
        self.assertEqual(service._date("26/06/27"), "2026-06-27")
        consumption = [{"id":"c1", "block_no":1, "source_row":2, "customer_code":"104370", "customer_material_code":"P1", "quantity":48, "price":1}]
        orders = [{"order_no":"341-1", "line_no":"1", "order_date":"2026-01-01", "customer_code":"104370", "customer_material_code":"P1", "factory_part_no":"6900026270", "quantity":48}]
        stock = [
            {"warehouse_code":"0087B", "factory_part_no":"6900026270", "lot_no":"CG7QBK542H4", "quantity":48, "stagnation_date":service._date("26/07/18")},
            {"warehouse_code":"0087B", "factory_part_no":"6900026270", "lot_no":"CG61AC515H4", "quantity":60, "stagnation_date":service._date("26/06/27")},
        ]
        config = {"104370": {"name":"南通深南", "warehouses":{"0087", "0087B"}}}
        confirmed_inventory = {("0087B", "6900026270", "CG61AC515H4"): 12}
        details, unmatched = service._match(consumption, orders, stock, config, {}, confirmed_inventory)
        self.assertEqual(unmatched, [])
        self.assertEqual([(row["lot_no"], row["quantity"]) for row in details], [("CG61AC515H4", 48)])

    def test_maps_multiple_inventory_uploads_by_warehouse_filename(self):
        files = [SimpleNamespace(filename=f"{warehouse}.xls") for warehouse in service.WAREHOUSES]
        mapped = service.normalize_inventory_uploads(files, service.WAREHOUSES)
        self.assertEqual(tuple(mapped), service.WAREHOUSES)
        with self.assertRaisesRegex(ValueError, "重复上传仓库"):
            service.normalize_inventory_uploads(files + [SimpleNamespace(filename="0079.xlsx")], service.WAREHOUSES)
        with self.assertRaisesRegex(ValueError, "未对应启用仓库"):
            service.normalize_inventory_uploads([*files[:-1], SimpleNamespace(filename="9999.xls")], service.WAREHOUSES)
        with self.assertRaisesRegex(ValueError, "请上传库存表"):
            service.normalize_inventory_uploads(files[:-1], service.WAREHOUSES)
        with self.assertRaisesRegex(ValueError, "仅支持"):
            service.normalize_inventory_uploads([*files[:-1], SimpleNamespace(filename="0087B.csv")], service.WAREHOUSES)
