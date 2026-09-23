from __future__ import annotations

import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from fangzheng_web_app import db
from fangzheng_web_app.mail_transcode_agent import mail_store
from fangzheng_web_app.order_entry_service import (
    QUOTE_SOURCE_PRICE_FIELDS,
    _initial_order_change_template_data,
    _initial_quote_groups,
    build_quote_template_export,
    get_order_change_template,
    get_quote_template,
    order_change_template_progress,
    quote_template_progress,
    calculate_quote_template_prices,
    quote_template_reconciliation,
    queue_order_change_template_extraction,
    queue_quote_template_extraction,
    reextract_order_change_template,
    run_template_extraction_task,
    save_order_change_template,
    save_quote_template,
)
from fangzheng_web_app.order_interface_service import (
    _flatten_order_info_candidates,
    _match_order_change_line,
    _store_order_change_matches,
    get_order_change_matches,
    get_order_detail_records,
    query_order_info,
    select_order_change_candidate,
    submit_aps_order_demand_import,
)
from fangzheng_web_app.order_intake_service import bootstrap_cases, list_cases


class OrderChangeTemplateMarkupTests(unittest.TestCase):
    def test_match_details_do_not_repeat_original_template_values(self) -> None:
        markup = (
            Path(__file__).resolve().parents[1]
            / "templates"
            / "order_automation_order_change_template.html"
        ).read_text(encoding="utf-8")

        self.assertNotIn("原模板明细", markup)
        self.assertNotIn("oc-original", markup)
        self.assertIn("ERP 匹配明细", markup)

    def test_erp_candidate_display_fields_hide_unneeded_headers(self) -> None:
        route_source = (
            Path(__file__).resolve().parents[1]
            / "fangzheng_web_app"
            / "routes.py"
        ).read_text(encoding="utf-8")

        self.assertIn('("scta01", "NYEOS订单号")', route_source)
        self.assertIn('("account_set", "账套")', route_source)
        self.assertIn('("sctb43", "厂别")', route_source)
        self.assertIn('("sctb35", "项次")', route_source)
        self.assertNotIn('("sctb04", "项次")', route_source)
        self.assertNotIn('("scta11", "送货客户ID")', route_source)
        self.assertNotIn('("scta38", "单头客户订单号")', route_source)
        self.assertNotIn('("peag04", "主数据品名规格")', route_source)

    def test_change_submit_modal_is_hidden_until_the_button_is_clicked(self) -> None:
        markup = (
            Path(__file__).resolve().parents[1]
            / "templates"
            / "order_automation_order_change_template.html"
        ).read_text(encoding="utf-8")

        self.assertIn('id="orderChangeModal" hidden', markup)
        self.assertIn('id="submitOrderChange" type="button" {% if template_locked %}disabled{% endif %}>提交APS</button>', markup)
        self.assertIn('.oc-modal-mask[hidden]{display:none}', markup)
        self.assertIn("open.addEventListener('click',()=>{modal.hidden=false", markup)

    def test_change_template_exposes_reextract_action(self) -> None:
        markup = (
            Path(__file__).resolve().parents[1]
            / "templates"
            / "order_automation_order_change_template.html"
        ).read_text(encoding="utf-8")

        self.assertIn("重新提取", markup)
        self.assertIn("order_automation_order_change_template_refresh", markup)

    def test_change_template_exposes_bulk_spreadsheet_controls(self) -> None:
        markup = (
            Path(__file__).resolve().parents[1]
            / "templates"
            / "order_automation_order_change_template.html"
        ).read_text(encoding="utf-8")

        self.assertIn('id="orderChangeBulkTools"', markup)
        self.assertIn('data-bulk-fill-down', markup)
        self.assertIn('data-bulk-edit', markup)
        self.assertIn("excludedFields: ['line_no']", markup)
        self.assertIn("{% if template_locked %}disabled{% endif %}", markup)

    def test_order_change_case_offers_reply_after_aps_submission(self) -> None:
        markup = (
            Path(__file__).resolve().parents[1]
            / "templates"
            / "order_automation_case.html"
        ).read_text(encoding="utf-8")

        self.assertIn("change_progress.stage == 'pending_reply'", markup)
        self.assertIn("查看修改订单", markup)

    def test_change_extraction_auto_open_refreshes_then_redirects_to_template(self) -> None:
        route_source = (
            Path(__file__).resolve().parents[1] / "fangzheng_web_app" / "routes.py"
        ).read_text(encoding="utf-8")
        markup = (
            Path(__file__).resolve().parents[1] / "templates" / "order_automation_case.html"
        ).read_text(encoding="utf-8")
        self.assertIn('query["auto_open_change"] = "1"', route_source)
        self.assertIn('"main.order_automation_order_change_template"', route_source)
        self.assertIn("change_progress.stage == 'extracting' and auto_open_change", markup)

    def test_order_overview_exposes_pending_reply_filter(self) -> None:
        route_source = (
            Path(__file__).resolve().parents[1]
            / "fangzheng_web_app"
            / "routes.py"
        ).read_text(encoding="utf-8")

        self.assertIn('"pending_reply": "待回复邮件"', route_source)

    def test_quote_template_reuses_erp_actions_without_aps_submission(self) -> None:
        markup = (
            Path(__file__).resolve().parents[1]
            / "templates"
            / "order_automation_order_change_template.html"
        ).read_text(encoding="utf-8")

        self.assertIn("is_quote_template", markup)
        self.assertIn("'核价模板' if is_quote_template", markup)
        self.assertIn("价格计算", markup)
        self.assertIn("can_calculate_quote_prices", markup)
        self.assertIn("当前邮件尚未关联客户", markup)
        self.assertIn("column_widths", markup)
        self.assertIn("oc-quote .oc-grid", markup)
        self.assertIn('id="quoteMatchModal" hidden', markup)
        self.assertIn('id="quoteReviewModal" hidden', markup)
        self.assertIn("oc-wide-modal", markup)
        self.assertIn("data-review-cell", markup)
        self.assertIn("panel?panel.cloneNode(true)", markup)
        self.assertNotIn("detail.cloneNode(true)", markup)
        self.assertIn("一键核对", markup)
        self.assertIn("导出核价表", markup)
        self.assertIn("order_automation_quote_template_export", markup)
        self.assertIn("order_automation_quote_template_query", markup)
        self.assertIn("order_automation_quote_template_select_match", markup)
        self.assertIn("oc-source-tabs", markup)
        self.assertIn("quote_groups", markup)
        self.assertIn("active_quote_group", markup)
        self.assertIn("{% if not is_quote_template %}<div class=\"oc-modal-mask\"", markup)

    def test_quote_template_uses_a_right_drawer_for_matched_erp_details(self) -> None:
        markup = (
            Path(__file__).resolve().parents[1]
            / "templates"
            / "order_automation_order_change_template.html"
        ).read_text(encoding="utf-8")

        self.assertIn('id="quoteMatchDrawer"', markup)
        self.assertIn('class="oc-quote-drawer"', markup)
        self.assertIn("item?.match_status==='matched'", markup)
        self.assertIn("ERP订单号", markup)
        self.assertIn("NYEOS订单号", markup)
        self.assertIn("客户订单号", markup)
        self.assertIn("客户与厂内差异", markup)
        self.assertIn("oc-review-tag", markup)

    def test_quote_template_posts_only_the_active_group_as_json(self) -> None:
        markup = (
            Path(__file__).resolve().parents[1]
            / "templates"
            / "order_automation_order_change_template.html"
        ).read_text(encoding="utf-8")

        self.assertIn("'Content-Type':'application/json'", markup)
        self.assertIn("group_key:groupKey,lines:readLines()", markup)
        self.assertNotIn("quoteGroups={{", markup)
        self.assertIn("data-quote-json-action=\"query\"", markup)


class OrderChangeTemplateTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db_patch = patch.object(db, "DATABASE_PATH", Path(self.temp_dir.name) / "change.sqlite3")
        self.db_patch.start()
        db.init_db()
        account_id = mail_store.create_or_update_account(
            "orders@example.com", owner_employee_id="employee-a", auth_code="auth-code"
        )
        mail_store.upsert_message(
            account_id, folder="INBOX", uid="change-1", message_id="<change-1@example.com>",
            subject="采购订单", sender="buyer@example.com", sent_at="2026-08-19 09:00:00",
            received_at="2026-08-19 09:00:00", body_html="", body_text="PO-CHANGE-001", eml_path="", is_order=1,
        )
        bootstrap_cases("employee-a", account_id)
        self.case_id = list_cases("employee-a", "2026-08-19", "new_order", account_id)[0]["id"]
        with db.db_cursor() as conn:
            conn.execute("UPDATE order_intake_cases SET action_type='order_change' WHERE id=?", (self.case_id,))

    def tearDown(self) -> None:
        self.db_patch.stop()
        self.temp_dir.cleanup()

    def test_change_template_extracts_in_background_and_saves_the_six_fields(self) -> None:
        with patch("fangzheng_web_app.order_entry_service.subprocess.Popen"):
            queued = queue_order_change_template_extraction(self.case_id, "employee-a")
        self.assertEqual(queued["status"], "queued")
        self.assertEqual(order_change_template_progress(self.case_id, "employee-a")["stage"], "extracting")

        header = {"customer_order_number": "PO-CHANGE-001"}
        lines = [{"values": {
            "line_no": "8", "customer_product_code": "CUST-8", "customer_spec": "NY2150 1.0mm",
            "delivery_date": "2026-09-25", "quantity": "100",
        }, "sources": {}}]
        with patch("fangzheng_web_app.order_entry_service.review_case_template_prices") as price_review, patch(
            "fangzheng_web_app.order_entry_service._initial_template_data", return_value=(header, lines)
        ):
            run_template_extraction_task(
                queued["task_id"], self.case_id, "employee-a", action_type="order_change",
            )
        price_review.assert_not_called()

        _case, extracted = get_order_change_template(self.case_id, "employee-a")
        self.assertEqual(extracted["header"]["customer_order_number"], "PO-CHANGE-001")
        self.assertEqual(extracted["lines"][0]["values"]["line_no"], "8")
        self.assertEqual(extracted["lines"][0]["values"]["customer_product_code"], "CUST-8")
        self.assertEqual(extracted["lines"][0]["values"]["unit_price"], "")
        extracted_details = get_order_detail_records(self.case_id, "employee-a")
        self.assertTrue(any(item["event_type"] == "order_change_template_extracted" for item in extracted_details["events"]))

        saved = save_order_change_template(self.case_id, "employee-a", {"lines": [{"values": {
            "customer_order_number": "PO-CHANGE-002", "line_no": "99", "customer_product_code": "CUST-9",
            "customer_spec": "NY2170 0.8mm", "delivery_date": "2026-10-01", "quantity": "80",
        }}]})
        values = saved["lines"][0]["values"]
        self.assertEqual(values["line_no"], "99")
        self.assertEqual(values["customer_order_number"], "PO-CHANGE-002")
        self.assertEqual(values["delivery_date"], "2026-10-01")
        self.assertEqual(order_change_template_progress(self.case_id, "employee-a")["stage"], "saved")
        saved_details = get_order_detail_records(self.case_id, "employee-a")
        self.assertTrue(any(item["event_type"] == "order_change_template_saved" for item in saved_details["events"]))

    def test_quote_template_calculates_and_invalidates_price_when_inputs_change(self) -> None:
        with db.db_cursor() as conn:
            conn.execute("UPDATE order_intake_cases SET action_type='quotation' WHERE id=?", (self.case_id,))
        with patch("fangzheng_web_app.order_entry_service.subprocess.Popen"):
            queued = queue_quote_template_extraction(self.case_id, "employee-a")
        with patch("fangzheng_web_app.order_entry_service._initial_template_data", return_value=(
            {"customer_order_number": "PO-QUOTE-1"},
            [{"values": {
                "line_no": "1", "customer_order_number": "PO-QUOTE-1", "customer_spec": "NY2150", "quantity": "10",
                "price_before_tax": "10.00", "amount_before_tax": "100.00", "unit_price": "11.30", "amount_with_tax": "113.00",
            }, "sources": {}}],
        )):
            run_template_extraction_task(queued["task_id"], self.case_id, "employee-a", action_type="quotation")
        self.assertEqual(quote_template_progress(self.case_id, "employee-a")["stage"], "pending_template_save")
        with patch("fangzheng_web_app.order_entry_service.review_case_template_prices", return_value={
            "association": {"matched": True, "price_customer_key": "fangzheng"},
            "by_line": {1: {"quote_price": "12.50", "note": "报价单计算"}},
        }):
            result = calculate_quote_template_prices(self.case_id, "employee-a")
        self.assertEqual(result["success"], 1)
        _case, quoted = get_quote_template(self.case_id, "employee-a")
        self.assertEqual(quoted["lines"][0]["values"]["quote_price"], "12.50")
        self.assertEqual(quoted["lines"][0]["values"]["quote_note"], "报价单计算")
        self.assertEqual(
            {field: quoted["lines"][0]["values"][field] for field in QUOTE_SOURCE_PRICE_FIELDS},
            {"price_before_tax": "10.00", "amount_before_tax": "100.00", "unit_price": "11.30", "amount_with_tax": "113.00"},
        )

        saved = save_quote_template(self.case_id, "employee-a", {"lines": [{"values": {
            "customer_order_number": "PO-QUOTE-1", "line_no": "1", "customer_spec": "NY2150", "quantity": "20",
            "price_before_tax": "10.00", "amount_before_tax": "200.00", "unit_price": "11.30", "amount_with_tax": "226.00",
        }}]})
        self.assertEqual(saved["lines"][0]["values"].get("quote_price"), "")
        self.assertEqual(saved["lines"][0]["values"].get("quote_note"), "")
        self.assertEqual(
            {field: saved["lines"][0]["values"][field] for field in QUOTE_SOURCE_PRICE_FIELDS},
            {"price_before_tax": "10.00", "amount_before_tax": "200.00", "unit_price": "11.30", "amount_with_tax": "226.00"},
        )

    def test_quote_template_partial_save_keeps_other_attachment_groups(self) -> None:
        with db.db_cursor() as conn:
            conn.execute("UPDATE order_intake_cases SET action_type='quotation' WHERE id=?", (self.case_id,))
        with patch("fangzheng_web_app.order_entry_service.subprocess.Popen"):
            queued = queue_quote_template_extraction(self.case_id, "employee-a")
        with patch("fangzheng_web_app.order_entry_service._initial_template_data", return_value=(
            {"customer_order_number": "PO-QUOTE-1"},
            [{"values": {"line_no": "1", "customer_order_number": "PO-QUOTE-1", "customer_spec": "初始"}, "sources": {}}],
        )):
            run_template_extraction_task(queued["task_id"], self.case_id, "employee-a", action_type="quotation")

        saved = save_quote_template(self.case_id, "employee-a", {"groups": [
            {"group_key": "attachment-one", "header": {"_quote_source_label": "F3.xlsx"}, "lines": [
                {"values": {"line_no": "1", "customer_order_number": "PO-1", "customer_spec": "规格一", "quantity": "10"}, "sources": {"source": "one"}},
            ]},
            {"group_key": "attachment-two", "header": {"_quote_source_label": "F5.xlsx"}, "lines": [
                {"values": {"line_no": "2", "customer_order_number": "PO-2", "customer_spec": "规格二", "quantity": "20"}, "sources": {"source": "two"}},
            ]},
        ]})
        partial = save_quote_template(self.case_id, "employee-a", {
            "group_key": "attachment-one",
            "lines": [{"values": {"line_no": "1", "customer_order_number": "PO-1", "customer_spec": "修改后规格", "quantity": "10"}}],
        })
        groups = {group["group_key"]: group for group in partial["groups"]}
        self.assertEqual(groups["attachment-one"]["lines"][0]["values"]["customer_spec"], "修改后规格")
        self.assertEqual(groups["attachment-two"]["lines"][0]["values"]["customer_spec"], "规格二")
        self.assertEqual(groups["attachment-two"]["lines"][0]["sources"], {"source": "two"})

    def test_order_query_batches_more_than_fifty_purchase_orders(self) -> None:
        calls: list[list[str]] = []

        def query_batch(*_args, **kwargs):
            numbers = list(_args[4])
            calls.append(numbers)
            return {
                "call_id": len(calls), "mode": "mock", "status": "success",
                "response": {"data": {"orderList": [], "notFoundList": []}},
            }

        orders = [f"PO-BATCH-{index:03d}" for index in range(101)]
        with patch("fangzheng_web_app.order_interface_service._query_order_info", side_effect=query_batch), patch(
            "fangzheng_web_app.order_interface_service._store_order_change_matches", return_value=[]
        ), patch("fangzheng_web_app.order_interface_service.record_order_detail_event"):
            result = query_order_info(self.case_id, 1, "employee-a", "employee-a", orders)

        self.assertEqual([len(batch) for batch in calls], [50, 50, 1])
        self.assertEqual(result["call_ids"], [1, 2, 3])
        self.assertEqual(result["matches"], [])

    def test_quote_template_groups_multiple_attachments_into_independent_tabs(self) -> None:
        header = {"customer_order_number": "PO-QUOTE-1"}
        lines = [
            {"values": {"line_no": "1", "customer_order_number": "PO-QUOTE-1", "customer_spec": "规格一"}, "sources": {
                "customer_spec": {"label": "附件：F3&F7改价明细.xlsx", "reference": "Sheet1 第 2 行"},
            }},
            {"values": {"line_no": "2", "customer_order_number": "PO-QUOTE-2", "customer_spec": "规格二"}, "sources": {
                "customer_spec": {"label": "附件：F5A 改价明细.pdf", "reference": "识别明细第 1 行"},
            }},
        ]

        groups = _initial_quote_groups(header, lines)

        self.assertEqual([group["header"]["_quote_source_label"] for group in groups], ["F3&F7改价明细.xlsx", "F5A 改价明细.pdf"])
        self.assertEqual([len(group["lines"]) for group in groups], [1, 1])

    def test_quote_template_keeps_an_empty_tab_for_an_unparsed_attachment(self) -> None:
        groups = _initial_quote_groups(
            {"customer_order_number": "PO-QUOTE-1"},
            [{"values": {"line_no": "1", "customer_spec": "规格一"}, "sources": {
                "customer_spec": {"label": "附件：已识别.xls", "reference": "Sheet1 第 2 行"},
            }}],
            source_labels=["未识别.xlsx", "已识别.xls"],
        )

        self.assertEqual([group["header"]["_quote_source_label"] for group in groups], ["未识别.xlsx", "已识别.xls"])
        self.assertEqual([len(group["lines"]) for group in groups], [0, 1])

    def test_quote_template_save_preserves_attachment_groups(self) -> None:
        with db.db_cursor() as conn:
            conn.execute("UPDATE order_intake_cases SET action_type='quotation' WHERE id=?", (self.case_id,))
        with patch("fangzheng_web_app.order_entry_service.subprocess.Popen"):
            queued = queue_quote_template_extraction(self.case_id, "employee-a")
        with patch("fangzheng_web_app.order_entry_service._initial_template_data", return_value=(
            {"customer_order_number": "PO-QUOTE-1"},
            [{"values": {"line_no": "1", "customer_order_number": "PO-QUOTE-1", "customer_spec": "初始"}, "sources": {}}],
        )):
            run_template_extraction_task(queued["task_id"], self.case_id, "employee-a", action_type="quotation")

        saved = save_quote_template(self.case_id, "employee-a", {"groups": [
            {"group_key": "attachment-one", "header": {"_quote_source_label": "F3&F7改价明细.xlsx"}, "lines": [
                {"values": {"line_no": "1", "customer_order_number": "PO-QUOTE-1", "customer_spec": "规格一", "quantity": "10"}},
            ]},
            {"group_key": "attachment-two", "header": {"_quote_source_label": "F5A 改价明细.pdf"}, "lines": [
                {"values": {"line_no": "2", "customer_order_number": "PO-QUOTE-2", "customer_spec": "规格二", "quantity": "20"}},
            ]},
        ]})

        self.assertEqual([group["display_order_number"] for group in saved["groups"]], ["F3&F7改价明细.xlsx", "F5A 改价明细.pdf"])
        self.assertEqual([[line["values"]["customer_spec"] for line in group["lines"]] for group in saved["groups"]], [["规格一"], ["规格二"]])

    def test_quote_template_export_uses_selected_erp_request_date(self) -> None:
        template = {
            "header": {"customer_order_number": "PO-EXPORT-001"},
            "lines": [
                {"line_no": 1, "values": {
                    "line_no": "1", "customer_product_code": "CUST-1", "customer_spec": "NY2150 1.0mm",
                    "delivery_date": "2026-10-01", "price_before_tax": "10.0000", "unit_price": "11.3000",
                    "quote_price": "12.5000",
                }},
                {"line_no": 2, "values": {
                    "line_no": "2", "customer_product_code": "CUST-2", "customer_spec": "NY2160 0.8mm",
                    "delivery_date": "2026-10-02", "price_before_tax": "20", "unit_price": "22.60",
                    "quote_price": "",
                }},
            ],
        }
        matches = {1: {"selected_candidate": {
            "scta39": "ERP-001", "account_set": "KL01", "sctb35": "9", "sctb16": "2026-11-08 00:00:00",
            "sctb07": "10.0000", "sctb06": "11.3000",
        }}}

        output, filename = build_quote_template_export(self.case_id, template, matches)
        workbook = load_workbook(BytesIO(output.getvalue()), data_only=True)
        sheet = workbook["核价表"]

        self.assertEqual(
            [cell.value for cell in sheet[1]],
            ["ERP订单号", "账套", "项次", "客户料号", "客户规格", "客户需求日期", "客户税前单价", "客户税后单价", "厂内税前单价", "厂内税后单价", "报价单价格"],
        )
        self.assertEqual([sheet.cell(2, column).value for column in range(1, 6)], ["ERP-001", "KL01", "9", "CUST-1", "NY2150 1.0mm"])
        self.assertEqual(sheet.cell(2, 6).value.date().isoformat(), "2026-11-08")
        self.assertEqual([sheet.cell(2, column).value for column in range(7, 12)], [10, 11.3, 10, 11.3, 12.5])
        self.assertEqual([sheet.cell(3, column).value for column in range(1, 4)], [None, None, None])
        self.assertEqual(sheet.cell(3, 6).value, None)
        self.assertIn("PO-EXPORT-001", filename)
        workbook.close()

    def test_quote_reconciliation_reports_price_and_amount_differences(self) -> None:
        template = {"lines": [{"line_no": 1, "values": {
            "line_no": "1", "quantity": "10", "price_before_tax": "10", "unit_price": "11.30",
            "amount_before_tax": "100", "amount_with_tax": "113", "quote_price": "11.30",
        }}]}
        matches = {1: {"status": "matched", "selected_candidate": {
            "sctb05": "10", "sctb07": "10", "sctb06": "11.30",
        }}}
        review = quote_template_reconciliation(template, matches, {"tax_mode": "tax_inclusive"})
        self.assertEqual(review["matched"], 1)
        self.assertEqual(review["issues"], 0)

        matches[1]["selected_candidate"]["sctb06"] = "12"
        review = quote_template_reconciliation(template, matches, {"tax_mode": "tax_inclusive"})
        self.assertEqual(review["issues"], 1)
        self.assertTrue(any(item["field"] == "含税单价" for item in review["rows"][0]["differences"]))

    def test_quote_reconciliation_does_not_round_unit_prices_before_comparing(self) -> None:
        template = {"lines": [{"line_no": 1, "values": {
            "line_no": "1", "quantity": "10", "price_before_tax": "10.0001", "unit_price": "11.3001",
            "amount_before_tax": "100", "amount_with_tax": "113", "quote_price": "",
        }}]}
        matches = {1: {"status": "matched", "selected_candidate": {
            "sctb05": "10", "sctb07": "10.0000", "sctb06": "11.3000",
        }}}

        review = quote_template_reconciliation(template, matches, {"tax_mode": "tax_inclusive"})

        differences = {item["field"]: item for item in review["rows"][0]["differences"]}
        self.assertEqual(review["issues"], 1)
        self.assertEqual(differences["未税单价"]["customer"], "10.0001")
        self.assertEqual(differences["未税单价"]["factory"], "10.0000")
        self.assertEqual(differences["含税单价"]["customer"], "11.3001")
        self.assertEqual(differences["含税单价"]["factory"], "11.3000")

    def test_quote_reconciliation_calculates_factory_amount_before_rounding_unit_price(self) -> None:
        template = {"lines": [{"line_no": 1, "values": {
            "line_no": "1", "quantity": "1300", "price_before_tax": "239.5200000",
            "amount_before_tax": "311376.00", "unit_price": "270.657600",
            "amount_with_tax": "351854.88", "quote_price": "239.52",
        }}]}
        matches = {1: {"status": "matched", "selected_candidate": {
            "sctb05": "1300", "sctb07": "239.52", "sctb06": "270.6576",
        }}}

        review = quote_template_reconciliation(template, matches, {"tax_mode": "tax_exclusive"})

        self.assertEqual(review["matched"], 1)
        self.assertEqual(review["issues"], 0)

    def test_quote_reconciliation_passes_when_quote_rule_does_not_match(self) -> None:
        template = {"lines": [{"line_no": 1, "values": {
            "line_no": "1", "quantity": "10", "price_before_tax": "10", "unit_price": "11.30",
            "amount_before_tax": "100", "amount_with_tax": "113", "quote_price": "",
            "quote_note": "报价未命中",
        }}]}
        matches = {1: {"status": "matched", "selected_candidate": {
            "sctb05": "10", "sctb07": "10", "sctb06": "11.30",
        }}}

        review = quote_template_reconciliation(template, matches, {"tax_mode": "tax_inclusive"})

        self.assertEqual(review["matched"], 1)
        self.assertEqual(review["issues"], 0)

    def test_quote_reconciliation_keeps_customer_quote_and_factory_prices_separate(self) -> None:
        template = {"lines": [{"line_no": 1, "values": {
            "line_no": "1", "quantity": "10", "price_before_tax": "10", "unit_price": "11.30",
            "amount_before_tax": "100", "amount_with_tax": "113", "quote_price": "12.50",
        }}]}
        matches = {1: {"status": "matched", "selected_candidate": {
            "sctb05": "10", "sctb07": "10", "sctb06": "11.30",
        }}}

        review = quote_template_reconciliation(template, matches, {"tax_mode": "tax_inclusive"})

        difference = next(item for item in review["rows"][0]["differences"] if item["field"] == "报价单价格")
        self.assertEqual(review["rows"][0]["quote_label"], "报价单含税单价")
        self.assertEqual(review["rows"][0]["quote_price"], "12.50")
        self.assertEqual(difference["customer"], "11.30")
        self.assertEqual(difference["quote"], "12.50")
        self.assertEqual(difference["factory"], "11.30")

    def test_quote_reconciliation_builds_four_row_drawer_data_with_tax_conversion(self) -> None:
        template = {"lines": [{"line_no": 1, "values": {
            "line_no": "1", "quantity": "10", "price_before_tax": "10.00", "unit_price": "11.30",
            "amount_before_tax": "100.00", "amount_with_tax": "113.00", "quote_price": "11.30",
        }}]}
        matches = {1: {"status": "matched", "selected_candidate": {
            "scta39": "ERP-1", "account_set": "KL01", "sctb35": "1", "scta01": "NYEOS-1",
            "sctb15": "PO-1", "sctb14": "CP-1", "sctb36": "客户规格", "sctb05": "10",
            "sctb07": "10.00", "sctb06": "11.30",
        }}}

        row = quote_template_reconciliation(template, matches, {"tax_mode": "tax_inclusive"})["rows"][0]

        self.assertEqual([item["label"] for item in row["comparison_rows"]], ["未税单价", "未税金额", "含税单价", "含税金额"])
        self.assertEqual(row["comparison_rows"][0]["quote"], "10.00")
        self.assertEqual(row["comparison_rows"][1]["quote"], "100.00")
        self.assertEqual(row["comparison_rows"][2]["quote"], "11.30")
        self.assertEqual(row["comparison_rows"][3]["quote"], "113.00")
        self.assertEqual(row["comparison_rows"][0]["difference"], "客户 − 厂内 = 0.00")
        self.assertEqual(row["erp"]["erp_order_number"], "ERP-1")
        self.assertEqual(row["erp"]["customer_order_number"], "PO-1")

    def test_quote_reconciliation_keeps_customer_factory_match_when_quote_is_missing(self) -> None:
        template = {"lines": [{"line_no": 1, "values": {
            "line_no": "1", "quantity": "10", "price_before_tax": "10", "unit_price": "11.30",
            "amount_before_tax": "100", "amount_with_tax": "113", "quote_price": "", "quote_note": "报价未命中",
        }}]}
        matches = {1: {"status": "matched", "selected_candidate": {
            "sctb05": "10", "sctb07": "10", "sctb06": "11.30",
        }}}

        row = quote_template_reconciliation(template, matches, {"tax_mode": "tax_inclusive"})["rows"][0]

        self.assertEqual(row["status"], "matched")
        self.assertIn("报价未命中", row["comparison_rows"][2]["quote"])

    def test_change_template_extracts_contract_code_description_and_quantity_from_mail_table(self) -> None:
        case = {
            "id": self.case_id,
            "customer_id": None,
            "received_at": "2026-09-16 09:00:00",
            "body_text": "",
            "body_html": """
                <table>
                  <tr><th>项次</th><th>合同号</th><th>编码</th><th>描述</th><th>数量</th><th>供应商</th><th>说明</th><th>备注</th></tr>
                  <tr><td>1</td><td>PO-26T1-005682</td><td>AA3780089110290001</td><td>NY-P6HF 0.089mm 1/1 21.42×24.42 H/H</td><td>28.000000</td><td>NANYA</td><td>更新数量</td><td>合同取消</td></tr>
                </table>
            """,
        }

        header, lines = _initial_order_change_template_data(case)

        self.assertEqual(header["customer_order_number"], "PO-26T1-005682")
        values = lines[0]["values"]
        self.assertEqual(values["line_no"], "1")
        self.assertEqual(values["customer_order_number"], "PO-26T1-005682")
        self.assertEqual(values["customer_product_code"], "AA3780089110290001")
        self.assertEqual(values["customer_spec"], "NY-P6HF 0.089mm 1/1 21.42×24.42 H/H")
        self.assertEqual(values["quantity"], "28.000000")

    def test_quote_template_extracts_tax_inclusive_and_exclusive_prices_and_amounts(self) -> None:
        case = {
            "id": self.case_id,
            "customer_id": None,
            "received_at": "2026-09-16 09:00:00",
            "body_text": "",
            "body_html": """
                <table>
                  <tr><th>项次</th><th>合同号</th><th>编码</th><th>描述</th><th>数量</th><th>未税单价</th><th>未税金额</th><th>含税单价</th><th>含税金额</th></tr>
                  <tr><td>1</td><td>PO-QUOTE-001</td><td>AA001</td><td>NY2150</td><td>12</td><td>140.98</td><td>1691.76</td><td>159.3074</td><td>1911.69</td></tr>
                </table>
            """,
        }

        _header, lines = _initial_order_change_template_data(case, include_pricing=True)

        values = lines[0]["values"]
        self.assertEqual(values["price_before_tax"], "140.98")
        self.assertEqual(values["amount_before_tax"], "1691.76")
        self.assertEqual(values["unit_price"], "159.3074")
        self.assertEqual(values["amount_with_tax"], "1911.69")

    def test_change_template_reextract_replaces_lines_and_keeps_a_history_snapshot(self) -> None:
        with patch("fangzheng_web_app.order_entry_service.subprocess.Popen"):
            queued = queue_order_change_template_extraction(self.case_id, "employee-a")
        with patch("fangzheng_web_app.order_entry_service._initial_template_data", return_value=(
            {"customer_order_number": "PO-OLD"},
            [{"values": {"line_no": "1", "customer_order_number": "PO-OLD", "customer_product_code": "OLD-CODE", "quantity": "10"}, "sources": {}}],
        )):
            run_template_extraction_task(queued["task_id"], self.case_id, "employee-a", action_type="order_change")

        with patch("fangzheng_web_app.order_entry_service._initial_order_change_template_data", return_value=(
            {"customer_order_number": "PO-NEW"},
            [{"values": {"line_no": "1", "customer_order_number": "PO-NEW", "customer_product_code": "NEW-CODE", "customer_spec": "新规格", "quantity": "20"}, "sources": {}}],
        )):
            result = reextract_order_change_template(self.case_id, "employee-a")

        self.assertEqual(result["previous_line_count"], 1)
        self.assertEqual(result["line_count"], 1)
        _case, template = get_order_change_template(self.case_id, "employee-a")
        self.assertEqual(template["lines"][0]["values"]["customer_product_code"], "NEW-CODE")
        details = get_order_detail_records(self.case_id, "employee-a")
        self.assertTrue(any(item["event_type"] == "order_change_template_reextracted" for item in details["events"]))

    def test_change_template_rejects_a_non_change_case(self) -> None:
        with db.db_cursor() as conn:
            conn.execute("UPDATE order_intake_cases SET action_type='new_order' WHERE id=?", (self.case_id,))
        with self.assertRaisesRegex(ValueError, "修改订单"):
            queue_order_change_template_extraction(self.case_id, "employee-a")

    def test_order_info_query_only_sends_customer_order_numbers(self) -> None:
        with patch("fangzheng_web_app.order_entry_service.subprocess.Popen"):
            queued = queue_order_change_template_extraction(self.case_id, "employee-a")
        with patch("fangzheng_web_app.order_entry_service._initial_template_data", return_value=(
            {"customer_order_number": "PO-DEFAULT"},
            [{"values": {
                "line_no": "1", "customer_order_number": "PO-QUERY-001",
                "customer_product_code": "CUST-QUERY-001", "quantity": "10",
            }, "sources": {}}],
        )):
            run_template_extraction_task(queued["task_id"], self.case_id, "employee-a", action_type="order_change")
        _case, template = get_order_change_template(self.case_id, "employee-a")
        result = query_order_info(
            self.case_id, int(template["id"]), "employee-a", "employee-a",
            ["PO-QUERY-001", "PO-QUERY-001", ""],
        )
        self.assertEqual(result["mode"], "mock")
        details = get_order_detail_records(self.case_id, "employee-a")
        call = details["calls"][0]
        self.assertEqual(call["interface_key"], "order_info_query")
        self.assertEqual(call["request"], {"orderNumberList": ["PO-QUERY-001"]})
        self.assertEqual(call["interface_label"], "查询订单信息")
        matches = get_order_change_matches(self.case_id, int(template["id"]), "employee-a")
        self.assertEqual(matches[1]["status"], "matched")
        self.assertEqual(matches[1]["selected_candidate"]["scta39"], "220-MOCK-0001")

        save_order_change_template(self.case_id, "employee-a", {"lines": [{"values": {
            "customer_order_number": "PO-QUERY-001", "line_no": "1",
            "customer_product_code": "", "quantity": "99",
        }}]})
        self.assertEqual(get_order_change_matches(self.case_id, int(template["id"]), "employee-a"), {})

    def test_order_change_matching_uses_priority_and_deduplicates_erp_rows(self) -> None:
        first = {
            "scta39": "ERP-1", "scta01": "ORDER-1", "sctb02": "PART-1", "sctb14": "Cust-A",
            "sctb15": "PO-A", "sctb35": 10, "sctb05": 100, "sctb43": "S1",
        }
        second = {
            "scta39": "ERP-2", "scta01": "ORDER-2", "sctb02": "PART-2", "sctb14": "CUST-A",
            "sctb15": "po-a", "sctb35": 20, "sctb05": 200,
        }
        payload = {"data": {"orderList": [
            {"scta01": "ORDER-1", "scta39": "ERP-1", "acsn": "NY02", "sctbList": [first, dict(first)]},
            {"scta01": "ORDER-2", "scta39": "ERP-2", "sctbList": [second]},
        ]}}
        candidates = _flatten_order_info_candidates(payload)
        self.assertEqual(len(candidates), 2)
        self.assertEqual(candidates[0]["account_set"], "KL02")
        self.assertEqual(candidates[0]["sctb43"], "S1")
        by_item = _match_order_change_line({
            "customer_order_number": " PO-A ", "customer_product_code": "cust-a",
            "line_no": "10", "quantity": "200.0",
        }, candidates)
        self.assertEqual(by_item["match_level"], "order_part_item")
        self.assertEqual(by_item["selected"]["scta39"], "ERP-1")
        by_quantity = _match_order_change_line({
            "customer_order_number": "PO-A", "customer_product_code": "CUST-A",
            "line_no": "99", "quantity": "200.0",
        }, candidates)
        self.assertEqual(by_quantity["match_level"], "order_part_quantity")
        self.assertEqual(by_quantity["selected"]["scta39"], "ERP-2")
        multiple = _match_order_change_line({
            "customer_order_number": "PO-A", "customer_product_code": "CUST-A",
            "line_no": "99", "quantity": "999",
        }, candidates)
        self.assertEqual(multiple["status"], "multiple")
        self.assertEqual(len(multiple["candidates"]), 2)

    def test_order_match_falls_back_to_po_spec_and_quantity_when_customer_part_is_wrong(self) -> None:
        candidates = [{
            "scta39": "ERP-SPEC-1", "sctb14": "正确客户料号", "sctb15": "PO-SPEC-1",
            "sctb36": "NY6180L 0.152mm H/H 41*49", "sctb05": "80", "sctb35": "1",
        }]

        matched = _match_order_change_line({
            "customer_order_number": "PO-SPEC-1", "customer_product_code": "1错误料号",
            "customer_spec": "NY6180L 0.152mm H/H 41*49", "line_no": "1", "quantity": "80",
        }, candidates)

        self.assertEqual(matched["status"], "matched")
        self.assertEqual(matched["match_level"], "order_spec_quantity")
        self.assertEqual(matched["selected"]["scta39"], "ERP-SPEC-1")

    def test_multiple_match_selection_is_persisted(self) -> None:
        with patch("fangzheng_web_app.order_entry_service.subprocess.Popen"):
            queued = queue_order_change_template_extraction(self.case_id, "employee-a")
        with patch("fangzheng_web_app.order_entry_service._initial_template_data", return_value=(
            {"customer_order_number": "PO-MULTI"},
            [{"values": {
                "line_no": "9", "customer_order_number": "PO-MULTI",
                "customer_product_code": "CUST-MULTI", "quantity": "99",
            }, "sources": {}}],
        )):
            run_template_extraction_task(queued["task_id"], self.case_id, "employee-a", action_type="order_change")
        _case, template = get_order_change_template(self.case_id, "employee-a")
        initial = query_order_info(
            self.case_id, int(template["id"]), "employee-a", "employee-a", ["PO-MULTI"],
        )
        response = {"data": {"orderList": [{
            "scta01": "ORDER-M", "scta39": "ERP-M",
            "sctbList": [
                {"sctb02": "PART-1", "sctb14": "CUST-MULTI", "sctb15": "PO-MULTI", "sctb35": 1, "sctb05": 1},
                {"sctb02": "PART-2", "sctb14": "CUST-MULTI", "sctb15": "PO-MULTI", "sctb35": 2, "sctb05": 2},
            ],
        }]}}
        with db.db_cursor() as conn:
            _store_order_change_matches(
                conn, case_id=self.case_id, template_id=int(template["id"]), employee_id="employee-a",
                call_id=initial["call_id"], response_payload=response,
            )
        pending = get_order_change_matches(self.case_id, int(template["id"]), "employee-a")[9]
        self.assertEqual(pending["status"], "multiple")
        selected = pending["candidates"][1]
        select_order_change_candidate(
            self.case_id, int(template["id"]), "employee-a", 9, selected["candidate_key"], "employee-a",
        )
        persisted = get_order_change_matches(self.case_id, int(template["id"]), "employee-a")[9]
        self.assertEqual(persisted["status"], "matched")
        self.assertEqual(persisted["selected_candidate"]["sctb02"], "PART-2")

    def test_pending_reply_change_cannot_be_saved_or_submitted_again(self) -> None:
        with patch("fangzheng_web_app.order_entry_service.subprocess.Popen"):
            queued = queue_order_change_template_extraction(self.case_id, "employee-a")
        with patch("fangzheng_web_app.order_entry_service._initial_template_data", return_value=(
            {"customer_order_number": "PO-LOCK"},
            [{"values": {
                "line_no": "1", "customer_order_number": "PO-LOCK",
                "customer_product_code": "CUST-LOCK", "quantity": "10",
            }, "sources": {}}],
        )):
            run_template_extraction_task(queued["task_id"], self.case_id, "employee-a", action_type="order_change")
        _case, template = get_order_change_template(self.case_id, "employee-a")
        with db.db_cursor() as conn:
            conn.execute("UPDATE order_intake_cases SET status='pending_reply' WHERE id=?", (self.case_id,))
        with self.assertRaisesRegex(ValueError, "不能重复提交"):
            submit_aps_order_demand_import(
                self.case_id, int(template["id"]), "employee-a", "employee-a", "交期变更",
            )
        with self.assertRaisesRegex(ValueError, "不能再修改模板"):
            save_order_change_template(self.case_id, "employee-a", {"lines": []})
