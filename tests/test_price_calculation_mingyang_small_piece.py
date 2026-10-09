import unittest
from dataclasses import replace

from openpyxl import Workbook

from fangzheng_web_app.price_calculation_extended import ExtPpRule, ExtRules, calculate_extended_spec
from fangzheng_web_app.price_calculation_service import process_price_workbook


class MingyangSmallPieceTests(unittest.TestCase):
    def setUp(self):
        self.row = ExtPpRule(8, "通用PP", "NY2170P", "2116", 58, 58, 200, 49.5, 110.2)

    def calculate(self, size, row=None):
        rules = ExtRules("mingyang", [row or self.row], [])
        return calculate_extended_spec("mingyang", f"PP NY2170P 2116 RC58% {size}", rules)

    def test_examples(self):
        result = self.calculate("523mm*417mm")
        self.assertEqual("成功", result.status)
        self.assertEqual(19.23, result.price)
        self.assertIn("总片数=1146", result.note)
        result = self.calculate("548*625mm", replace(self.row, length=150, price=11840 / 150))
        self.assertEqual(21.68, result.price)
        self.assertIn("总片数=546", result.note)

    def test_fixed_orientation(self):
        result = self.calculate("417mm*523mm")
        self.assertIn("纬向开片数=2，经向开片数=479", result.note)
        self.assertEqual(23.01, result.price)

    def test_exact_division_and_unrounded_meter_price(self):
        result = self.calculate("500mm*419.1mm", replace(self.row, price=110.23456))
        self.assertIn("总片数=1200", result.note)
        self.assertEqual(18.37, result.price)

    def test_missing_or_invalid_quote_values(self):
        for field in ("width", "length", "price"):
            for value in (None, 0, -1):
                with self.subTest(field=field, value=value):
                    self.assertEqual("失败", self.calculate("417*523mm", replace(self.row, **{field: value})).status)

    def test_invalid_dimensions(self):
        for size in ("0*523mm", "-417*523mm", "417*0mm", "523*1300mm", "200001*417mm"):
            with self.subTest(size=size):
                self.assertEqual("失败", self.calculate(size).status)

    def test_regular_pp_unchanged(self):
        result = self.calculate("200M/卷")
        self.assertEqual(110.2, result.price)
        self.assertEqual(22040, result.total)

    def test_batch_matches_single(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["客户规格"])
        ws.append(["PP NY2170P 2116 RC58% 417mm*523mm"])
        process_price_workbook(wb, "mingyang", ExtRules("mingyang", [self.row], []))
        headers = {cell.value: cell.column for cell in ws[1]}
        self.assertEqual(self.calculate("417mm*523mm").price, ws.cell(2, headers["新价格"]).value)
        self.assertEqual("", ws.cell(2, headers["整卷价格"]).value)


if __name__ == "__main__":
    unittest.main()
