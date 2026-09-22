from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from fangzheng_web_app import hushi_service
from fangzheng_web_app.hushi_service import calculate_hushi_workbook


class HushiSpecParsingTests(unittest.TestCase):
    def test_accepts_hushi_pp_glass_and_decimal_quote_rc(self) -> None:
        self.assertEqual(
            ("1067", 73.0),
            hushi_service.extract_pp_fields("PP NY6300SLP 1067 RC73% 518X619mm SNY"),
        )
        self.assertTrue(hushi_service.percentage_match(73, 0.73))

    def test_accepts_unquoted_thickness_compound_stack_and_foil_alias(self) -> None:
        fields = hushi_service.extract_ccl_fields(
            "NY6300(C) 0050 1/1 510*485 1078*2 H-VIP JNY 配切"
        )

        self.assertEqual(5.0, fields["thickness"])
        self.assertEqual((("1078", 2),), fields["structure"])
        self.assertEqual("HVLP", fields["foil"])
        self.assertTrue(hushi_service.structure_match(
            hushi_service.normalize_ccl_structure("7628*2+2116*1"),
            "7628X2+2116X1",
        ))
        self.assertTrue(hushi_service.foil_match("STD", "SDT"))
        self.assertEqual("HTE", hushi_service.extract_ccl_fields(
            "NY3170M 0040 H/1 713X408 2113X1 HTE SNY"
        )["foil"])


class HushiWorkbookExportTests(unittest.TestCase):
    def test_exports_normal_rebate_and_rebate_final_price(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            rule_dir = root / "rules"
            rule_dir.mkdir()
            self._write_pp_quote(rule_dir / "NY2170.xlsx")

            input_path = root / "input.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["客户规格"])
            sheet.append(["PP NY2170 2116 RC58% 1000X1000"])
            sheet.append(["PP NY2170 2116 RC59% 1000X1000"])
            workbook.save(input_path)

            output_path, _total, success, failed, _skipped = calculate_hushi_workbook(input_path, rule_dir)
            output = load_workbook(output_path, data_only=True)
            sheet = output.worksheets[0]

            self.assertEqual(("Normal", "Rebate", "Normal+Rebate", "Rebate最终价格"), (
                sheet.cell(1, 2).value,
                sheet.cell(1, 3).value,
                sheet.cell(1, 4).value,
                sheet.cell(1, 5).value,
            ))
            self.assertEqual((10, 8, 10, 84.48), (
                sheet.cell(2, 2).value,
                sheet.cell(2, 3).value,
                sheet.cell(2, 4).value,
                sheet.cell(2, 5).value,
            ))
            self.assertEqual((12, 9, 12, 95.04), (
                sheet.cell(3, 2).value,
                sheet.cell(3, 3).value,
                sheet.cell(3, 4).value,
                sheet.cell(3, 5).value,
            ))
            self.assertEqual((2, 0), (success, failed))

    def test_uses_sqf_when_special_price_is_struck_through(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            rule_dir = root / "rules"
            rule_dir.mkdir()
            self._write_pp_quote(rule_dir / "NY2170.xlsx")

            input_path = root / "input.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["客户规格"])
            sheet.append(["PP NY2170 2116 RC60% 1000X1000"])
            sheet.append(["PP NY2170 2116 RC61% 1000X1000"])
            workbook.save(input_path)

            output_path, _total, success, failed, _skipped = calculate_hushi_workbook(input_path, rule_dir)
            output = load_workbook(output_path, data_only=True)
            sheet = output.worksheets[0]

            self.assertEqual((15, 11, 15, 158.4), (
                sheet.cell(2, 2).value,
                sheet.cell(2, 3).value,
                sheet.cell(2, 4).value,
                sheet.cell(2, 5).value,
            ))
            self.assertEqual((None, 13, 13, 137.28), (
                sheet.cell(3, 2).value,
                sheet.cell(3, 3).value,
                sheet.cell(3, 4).value,
                sheet.cell(3, 5).value,
            ))
            self.assertEqual((2, 0), (success, failed))

    def test_reuses_quote_row_index_for_batch_calculation(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            rule_dir = root / "rules"
            rule_dir.mkdir()
            quote_path = rule_dir / "NY2170.xlsx"
            self._write_pp_quote(quote_path)
            rule_cache = {"NY2170": quote_path}
            quote_cache = {}

            with patch.object(
                hushi_service,
                "iter_data_rows",
                wraps=hushi_service.iter_data_rows,
            ) as iter_rows:
                for rc in (58, 59, 60, 61):
                    result = hushi_service.calculate_hushi_spec(
                        f"PP NY2170 2116 RC{rc}% 1000X1000",
                        rule_dir,
                        rule_cache=rule_cache,
                        quote_cache=quote_cache,
                    )
                    self.assertEqual("completed", result.status)

            self.assertEqual(1, iter_rows.call_count)

    @staticmethod
    def _write_pp_quote(path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "PP报价"
        sheet.append(["产品", "玻纤", "RC", "SQF", "特殊价格"])
        sheet.append(["NY2170", "2116", 58, 10, 8])
        sheet.append(["NY2170", "2116", 59, 12, 9])
        sheet.append(["NY2170", "2116", 60, 15, 11])
        sheet.append(["NY2170", "2116", 61, None, 13])
        sheet.cell(4, 5).font = Font(strike=True)
        workbook.save(path)
