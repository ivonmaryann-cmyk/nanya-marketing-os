from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from fangzheng_web_app.in_transit_service import _write_customer_sheet, load_customer_rows


class InTransitCustomerHeaderTests(unittest.TestCase):
    def test_legacy_customer_headers_are_normalized_for_matching_and_output(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source_path = Path(directory) / "legacy_customer_detail.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "客户明细"
            worksheet.append([
                "订单创建日期", "业务实体名", "采购订单号", "订单行号", "需求日期", "承诺日期",
                "物料编码", "物料描述", "单位", "PO发运行数量",
            ])
            worksheet.append([
                "2026-09-01", "南通N5", "PO-001", "2", "2026-09-08", "2026-09-10",
                "MAT-001", "覆铜板示例", "张", 12,
            ])
            workbook.save(source_path)

            rows, _stats = load_customer_rows(source_path, sheet_name="客户明细")

            self.assertEqual(len(rows), 1)
            row = rows[0]
            self.assertEqual(row.key, ("PO-001", "MAT-001", "2", "2026-09-10"))
            self.assertEqual(list(row.values), [
                "PO创建日期", "业务实体名", "单据编号", "单据行号", "订单需求日期", "订单承诺日期",
                "物料编码", "物料说明", "物料单位", "物料数量",
            ])

            output = Workbook()
            output_sheet = output.active
            _write_customer_sheet(output_sheet, rows)
            output_path = Path(directory) / "result.xlsx"
            output.save(output_path)

            headers = [cell.value for cell in load_workbook(output_path).active[1]]
            self.assertEqual(headers[:11], [
                "PO创建日期", "业务实体名", "单据编号", "单据行号", "订单需求日期", "订单承诺日期",
                "物料编码", "物料说明", "品名", "物料单位", "物料数量",
            ])
