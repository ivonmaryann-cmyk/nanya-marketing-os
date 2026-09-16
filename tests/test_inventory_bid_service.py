from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from fangzheng_web_app.inventory_bid_service import build_inventory_bid_rows


HEADERS = ["规格", "类别", "厚度", "铜箔", "尺寸", "水印", "等级", "数量", "单重"]


def _append_row(sheet, *, grade: str, quantity: int, spec: str = "NY2150") -> None:
    sheet.append([spec, "NY", 0.5, "1/1", "M3", "HTE", grade, quantity, 1.2])


def _add_sheet(workbook: Workbook, title: str):
    sheet = workbook.active if len(workbook.worksheets) == 1 and workbook.active.max_row == 1 else workbook.create_sheet()
    sheet.title = title
    sheet.append(HEADERS)
    return sheet


class InventoryBidServiceTests(unittest.TestCase):
    def test_prefers_the_b3_only_sheet_over_a_mixed_first_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            shanghai_path = directory_path / "shanghai.xlsx"
            jiangxi_path = directory_path / "jiangxi.xlsx"

            shanghai = Workbook()
            mixed = _add_sheet(shanghai, "全部库存")
            _append_row(mixed, grade="A1", quantity=100)
            _append_row(mixed, grade="B3", quantity=3)
            b3_only = _add_sheet(shanghai, "B级库存")
            _append_row(b3_only, grade="B3", quantity=5)
            shanghai.save(shanghai_path)

            jiangxi = Workbook()
            jiangxi_sheet = _add_sheet(jiangxi, "B级库存")
            _append_row(jiangxi_sheet, grade="B3", quantity=7)
            jiangxi.save(jiangxi_path)

            rows, stats = build_inventory_bid_rows(shanghai_path, jiangxi_path)

            self.assertEqual(stats["read"], 2)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["上海"], 5)
            self.assertEqual(rows[0]["江西"], 7)
            self.assertEqual(rows[0]["总计"], 12)

    def test_filters_to_b3_when_only_a_mixed_sheet_is_available(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            shanghai_path = directory_path / "shanghai.xlsx"
            jiangxi_path = directory_path / "jiangxi.xlsx"

            shanghai = Workbook()
            mixed = _add_sheet(shanghai, "库存")
            _append_row(mixed, grade="A1", quantity=100)
            _append_row(mixed, grade="B3", quantity=5)
            shanghai.save(shanghai_path)

            jiangxi = Workbook()
            jiangxi_sheet = _add_sheet(jiangxi, "库存")
            _append_row(jiangxi_sheet, grade="B3", quantity=7)
            jiangxi.save(jiangxi_path)

            rows, stats = build_inventory_bid_rows(shanghai_path, jiangxi_path)

            self.assertEqual(stats["read"], 2)
            self.assertEqual(rows[0]["上海"], 5)
            self.assertEqual(rows[0]["总计"], 12)

    def test_rejects_a_file_without_b3_inventory(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            shanghai_path = directory_path / "shanghai.xlsx"
            jiangxi_path = directory_path / "jiangxi.xlsx"

            shanghai = Workbook()
            _append_row(_add_sheet(shanghai, "库存"), grade="A1", quantity=100)
            shanghai.save(shanghai_path)

            jiangxi = Workbook()
            _append_row(_add_sheet(jiangxi, "库存"), grade="B3", quantity=7)
            jiangxi.save(jiangxi_path)

            with self.assertRaisesRegex(ValueError, "未找到等级为 B3"):
                build_inventory_bid_rows(shanghai_path, jiangxi_path)
